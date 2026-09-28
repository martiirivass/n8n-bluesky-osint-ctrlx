"""
Genera las planillas de Excel para que los dos evaluadores clasifiquen la muestra, a partir de
validacion_muestra_80.csv (que produce sortear_muestra_validacion.py).

Por qué Excel y no el CSV: el CSV está en UTF-8 sin BOM, y Excel al abrirlo con doble click rompe
los acentos y los emojis. Las planillas .xlsx los muestran bien y además traen menús desplegables
que impiden errores de tipeo en las columnas de clasificación.

Se generan DOS planillas para que cada evaluador trabaje de forma INDEPENDIENTE (sin ver las
respuestas del otro; si se ven, el kappa deja de medir acuerdo real):
  - planilla_evaluador_A.xlsx : las 80 filas; completa relevancia temática, país, motivo y notas.
  - planilla_evaluador_B.xlsx : solo las primeras N_DOBLE filas (la muestra ya viene en orden
                                 aleatorio, así que son un subconjunto al azar); completa
                                 relevancia temática y país.
Para evitar anclaje, las planillas NO muestran keyword_busqueda ni el flag de New Mexico.

Uso:
    python3 validacion/preparar_planillas.py
    python3 validacion/preparar_planillas.py --a Enzo --b Martiniano --n-doble 40
Después de clasificar: python3 validacion/unir_clasificaciones.py
"""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

DIR = Path(__file__).resolve().parent
CSV_BASE = DIR / "validacion_muestra_80.csv"

CLASES = "relevante,parcial,no_relevante,no_clasificable"
GEO = "si,no,no_se"
MOTIVOS = "seguridad_publica,homonimo_geografico,ruido_linguistico,spam_bot,otro"

RUBRICA = [
    ("RELEVANCIA TEMÁTICA (columna 'clasificacion')", ""),
    ("relevante", "El post trata de ciberseguridad de forma sustantiva: incidentes, ataques, "
                  "amenazas, vulnerabilidades, normativa, defensa o concientización."),
    ("parcial", "La ciberseguridad aparece solo de pasada, o el foco del post es otro tema "
                "(un evento, un negocio, una noticia general) y la ciberseguridad es secundaria."),
    ("no_relevante", "No trata de ciberseguridad, aunque contenga la palabra buscada "
                     "(por ejemplo 'hackear' en sentido coloquial, política, seguridad pública)."),
    ("no_clasificable", "No hay texto que juzgar, o no se entiende y no se puede decidir."),
    ("", ""),
    ("RELEVANCIA GEOGRÁFICA (columna 'pais_correcto')", "Se responde en TODAS las filas clasificables, "
                                                        "aunque el post no sea relevante en lo temático. Es independiente de la columna anterior."),
    ("si", "El post trata del país indicado en 'pais_buscado': hechos ocurridos allí, organizaciones "
           "o personas de ese país, normativa o instituciones de ese país."),
    ("no", "El país solo aparece de pasada, o el post trata de otro lugar (otro país, o un homónimo "
           "como New Mexico)."),
    ("no_se", "El texto no alcanza para decidir."),
    ("", ""),
    ("MOTIVO (solo evaluador A)", "Si la fila es no_relevante o parcial, elegí el motivo principal: "
                                  "seguridad_publica, homonimo_geografico, ruido_linguistico, spam_bot u otro."),
    ("", ""),
    ("REGLAS", "1) Cada evaluador trabaja solo, sin mirar la planilla del otro. "
               "2) Las definiciones de arriba son un punto de partida: si las cambian, cámbienlas ANTES de "
               "empezar y que sean idénticas para los dos. "
               "3) No edites las columnas grises."),
]


def crear_planilla(ruta, filas, columnas_edit, evaluador, titulo_eval):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clasificar"

    encabezado = ["n", "post_uri", "texto", "pais_buscado"] + [c for c, _ in columnas_edit] + ["evaluador"]
    ws.append(encabezado)
    gris = PatternFill("solid", fgColor="E7E7E7")
    azul = PatternFill("solid", fgColor="DCEBFA")
    for i, cell in enumerate(ws[1], start=1):
        cell.font = Font(bold=True)
        cell.fill = gris if i <= 4 else azul
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for n, r in enumerate(filas, start=1):
        ws.append([n, r["post_uri"], r["texto"], r["pais"]] + [""] * len(columnas_edit) + [evaluador])
        for c in range(1, 5):
            ws.cell(row=n + 1, column=c).fill = gris
        ws.cell(row=n + 1, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        for c in range(1, 5):
            if c != 3:
                ws.cell(row=n + 1, column=c).alignment = Alignment(vertical="top")

    anchos = {"n": 5, "post_uri": 14, "texto": 90, "pais_buscado": 13}
    for i, nombre in enumerate(encabezado, start=1):
        letra = ws.cell(row=1, column=i).column_letter
        if nombre in anchos:
            ws.column_dimensions[letra].width = anchos[nombre]
        elif nombre == "notas":
            ws.column_dimensions[letra].width = 40
        else:
            ws.column_dimensions[letra].width = 20
    ws.column_dimensions["B"].hidden = False
    ws.freeze_panes = "D2"

    ultima = len(filas) + 1
    for idx, (nombre, lista) in enumerate(columnas_edit, start=5):
        if lista is None:
            continue
        letra = ws.cell(row=1, column=idx).column_letter
        estricto = nombre != "motivo_si_no_relevante"
        dv = DataValidation(type="list", formula1=f'"{lista}"', allow_blank=True,
                            showErrorMessage=True, errorStyle="stop" if estricto else "warning",
                            errorTitle="Valor no permitido",
                            error="Elegí un valor de la lista desplegable.")
        ws.add_data_validation(dv)
        dv.add(f"{letra}2:{letra}{ultima}")

    guia = wb.create_sheet("Instrucciones", 0)
    guia.append([f"Planilla de validación — {titulo_eval}"])
    guia["A1"].font = Font(bold=True, size=13)
    guia.append([])
    for a, b in RUBRICA:
        guia.append([a, b])
    guia.column_dimensions["A"].width = 48
    guia.column_dimensions["B"].width = 110
    for row in guia.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.active = 1  # abre directo en la hoja "Clasificar"
    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--a", default="Enzo", help="nombre del evaluador A")
    ap.add_argument("--b", default="Martiniano", help="nombre del evaluador B")
    ap.add_argument("--n-doble", type=int, default=40, help="filas que clasifica también el evaluador B")
    args = ap.parse_args()

    with open(CSV_BASE, encoding="utf-8", newline="") as f:
        filas = list(csv.DictReader(f))
    print(f"Muestra base: {len(filas)} filas")

    wa = crear_planilla(
        DIR / "planilla_evaluador_A.xlsx", filas,
        [("clasificacion", CLASES), ("pais_correcto", GEO), ("motivo_si_no_relevante", MOTIVOS), ("notas", None)],
        args.a, f"Evaluador A ({args.a}) — las {len(filas)} filas")
    wa.save(DIR / "planilla_evaluador_A.xlsx")

    wb = crear_planilla(
        DIR / "planilla_evaluador_B.xlsx", filas[: args.n_doble],
        [("clasificacion_b", CLASES), ("pais_correcto_b", GEO)],
        args.b, f"Evaluador B ({args.b}) — las primeras {args.n_doble} filas")
    wb.save(DIR / "planilla_evaluador_B.xlsx")

    print("Escrito: planilla_evaluador_A.xlsx (80 filas) y planilla_evaluador_B.xlsx "
          f"({args.n_doble} filas) en {DIR}")


if __name__ == "__main__":
    main()
